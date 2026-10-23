"""Build Wriggle_Fit_Design_Proposal.docx from the Markdown source (headings, paragraphs,
bullets, pipe tables, figures with numbered captions).

    PYTHONPATH=<dir with python-docx> ~/.venv/bin/python Tanvir/report/build_docx.py
"""
import re
import sys
from pathlib import Path

sys.dont_write_bytecode = True
from docx import Document                                   # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT              # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH               # noqa: E402
from docx.oxml import OxmlElement                           # noqa: E402
from docx.oxml.ns import qn                                 # noqa: E402
from docx.shared import Inches, Pt, RGBColor                # noqa: E402

HERE = Path(__file__).resolve().parent
SRC = HERE / "Wriggle_Fit_Design_Proposal.md"
OUT = HERE / "Wriggle_Fit_Design_Proposal.docx"
TEXT_W = 6.9          # inches of usable width (A4, 0.7 in margins)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def runs(par, text, size=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 10) - 1)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        if bold:
            r.bold = True
        if size and not part.startswith("`"):
            r.font.size = Pt(size)


def shade(cell, hexcol):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol)
    tcPr.append(sh)


def add_table(doc, rows):
    hdr, body = rows[0], rows[1:]
    n = len(hdr)
    t = doc.add_table(rows=1 + len(body), cols=n)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fs = 7.5 if n >= 6 else 8
    for j, h in enumerate(hdr):
        c = t.rows[0].cells[j]; c.text = ""
        runs(c.paragraphs[0], h, size=fs, bold=True); shade(c, "E5E7EB")
    for i, r in enumerate(body, start=1):
        for j in range(n):
            c = t.rows[i].cells[j]; c.text = ""
            runs(c.paragraphs[0], r[j] if j < len(r) else "", size=fs)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    for a in ("left_margin", "right_margin"):
        setattr(sec, a, Inches(0.7))
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.0
    for name, sz in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 10.5)):
        h = doc.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        h.paragraph_format.space_before = Pt(8); h.paragraph_format.space_after = Pt(3)
    lines = SRC.read_text().splitlines()
    fig_n, i = 0, 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("# "):
            p = doc.add_paragraph(); r = p.add_run(ln[2:]); r.bold = True; r.font.size = Pt(18)
            p.paragraph_format.space_after = Pt(2)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:], level=3)
        elif ln.startswith("## "):
            doc.add_heading(ln[3:], level=1)
        elif ln.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*)\)", ln)
            cap, path = m.group(1), HERE / m.group(2)
            fig_n += 1
            doc.add_picture(str(path), width=Inches(TEXT_W if "fig01" in path.name or "fig03" in path.name or "fig04" in path.name or "fig08" in path.name or "fig09" in path.name else 6.4))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            doc.paragraphs[-1].paragraph_format.space_after = Pt(1)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run("Figure %d. " % fig_n); r.bold = True; r.font.size = Pt(8.5)
            runs(p, cap, size=8.5)
            p.paragraph_format.space_after = Pt(6)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-| ]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows); continue
        elif re.match(r"^- ", ln):
            p = doc.add_paragraph(style="List Bullet"); runs(p, ln[2:]); p.paragraph_format.space_after = Pt(1)
        elif re.match(r"^\d+\. ", ln):
            p = doc.add_paragraph(style="List Number"); runs(p, re.sub(r"^\d+\. ", "", ln)); p.paragraph_format.space_after = Pt(1)
        else:
            p = doc.add_paragraph(); runs(p, ln)
        i += 1
    doc.save(OUT)
    print("saved", OUT, "figures:", fig_n)


if __name__ == "__main__":
    main()
